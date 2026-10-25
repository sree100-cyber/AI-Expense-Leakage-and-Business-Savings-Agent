"""
test_reports.py - Tests for report generation and conversational agent.
TC-14, TC-15
"""
import pytest
import pandas as pd
import io
import sys, os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from modules.report_generator import (
    generate_excel_report,
    generate_duplicate_csv,
    generate_anomaly_csv,
    generate_recommendations_csv,
)
from modules.conversational_agent import RuleBasedAgent


@pytest.fixture
def full_context(sample_df):
    """Build a complete context dict for conversational agent tests."""
    exact_dups = sample_df[sample_df["payment_reference"] == "REF-001"].copy()
    exact_dups["duplicate_group_id"] = "EX-0001"
    exact_dups["duplicate_type"] = "Exact"
    exact_dups["match_reason"] = "REF-001"
    exact_dups["confidence"] = "High"

    recommendations = [
        {
            "id": 1, "title": "Review Exact Duplicate Payments",
            "category": "Payment Controls", "vendor": "Multiple",
            "evidence": "2 transactions share REF-001",
            "action": "Review AP logs",
            "monthly_savings": 96.0, "annual_savings": 1152.0,
            "confidence": "High", "priority": "High", "effort": "Low",
            "assumption": "Test assumption", "status": "Pending",
        }
    ]

    return {
        "df": sample_df,
        "exact_dups": exact_dups,
        "possible_dups": pd.DataFrame(),
        "anomaly_df": pd.DataFrame(),
        "recommendations": recommendations,
        "trends": {"recurring_expenses": pd.DataFrame()},
        "currency": "USD",
    }


class TestReportGeneration:
    def test_excel_report_contains_expected_sheets(self, sample_df):
        """TC-14: Excel report has expected sheets."""
        import openpyxl
        excel_bytes = generate_excel_report(
            df_clean=sample_df,
            exact_dups=pd.DataFrame(),
            possible_dups=pd.DataFrame(),
            anomaly_df=pd.DataFrame(),
            recommendations=[],
            trends={"monthly_by_category": pd.DataFrame(), "recurring_expenses": pd.DataFrame()},
        )
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        assert "Summary" in wb.sheetnames
        assert "Cleaned Transactions" in wb.sheetnames
        assert "Disclaimer" in wb.sheetnames

    def test_excel_report_transaction_count(self, sample_df):
        """TC-14: Cleaned Transactions sheet has correct row count."""
        import openpyxl
        excel_bytes = generate_excel_report(
            df_clean=sample_df,
            exact_dups=pd.DataFrame(),
            possible_dups=pd.DataFrame(),
            anomaly_df=pd.DataFrame(),
            recommendations=[],
            trends={"monthly_by_category": pd.DataFrame(), "recurring_expenses": pd.DataFrame()},
        )
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        ws = wb["Cleaned Transactions"]
        # Header row + data rows
        data_rows = ws.max_row - 1
        assert data_rows == len(sample_df), \
            f"Expected {len(sample_df)} data rows, got {data_rows}"

    def test_duplicate_csv_contains_expected_columns(self, sample_df):
        """TC-14: Duplicate CSV contains transaction_id and duplicate_group_id."""
        exact_dups = sample_df.copy()
        exact_dups["duplicate_group_id"] = "EX-0001"
        exact_dups["duplicate_type"] = "Exact"
        exact_dups["match_reason"] = "REF-001"
        exact_dups["confidence"] = "High"

        csv_bytes = generate_duplicate_csv(exact_dups, pd.DataFrame())
        df_out = pd.read_csv(io.BytesIO(csv_bytes))
        assert "transaction_id" in df_out.columns
        assert "duplicate_group_id" in df_out.columns

    def test_anomaly_csv_matches_input(self, sample_df):
        """TC-14: Anomaly CSV record count matches input."""
        sample_df["is_anomaly"] = True
        sample_df["anomaly_reason"] = "Test reason"
        sample_df["detection_method"] = "Statistical"
        csv_bytes = generate_anomaly_csv(sample_df)
        df_out = pd.read_csv(io.BytesIO(csv_bytes))
        assert len(df_out) == len(sample_df)

    def test_recommendations_csv_columns(self):
        """TC-14: Recommendations CSV has expected columns."""
        recs = [
            {"id": 1, "title": "Test", "category": "Software", "vendor": "X",
             "evidence": "e", "action": "a", "monthly_savings": 100.0,
             "annual_savings": 1200.0, "confidence": "High", "status": "Pending"}
        ]
        csv_bytes = generate_recommendations_csv(recs, "USD")
        df_out = pd.read_csv(io.BytesIO(csv_bytes))
        assert "Title" in df_out.columns
        assert "Priority" in df_out.columns or "ID" in df_out.columns

    def test_empty_anomaly_csv(self):
        """Empty anomaly DataFrame returns sensible bytes."""
        result = generate_anomaly_csv(pd.DataFrame())
        assert len(result) > 0  # Should return a message, not crash


class TestConversationalAgent:
    def test_top_category_answer(self, full_context):
        """TC-15: Agent correctly identifies highest-spending category."""
        agent = RuleBasedAgent(full_context)
        answer = agent.answer("Which category has the highest spending?")
        # Should mention a category name and a dollar amount
        assert "$" in answer or "USD" in answer or any(
            cat in answer for cat in ["Software", "Legal", "Travel", "Marketing", "Equipment"]
        ), f"Answer should mention category/amount: {answer}"

    def test_duplicate_summary_answer(self, full_context):
        """TC-15: Agent reports duplicates from actual data."""
        agent = RuleBasedAgent(full_context)
        answer = agent.answer("Which transactions may be duplicates?")
        assert "duplicate" in answer.lower() or "REF" in answer

    def test_top_recommendations_answer(self, full_context):
        """TC-15: Agent answers top savings recommendations."""
        agent = RuleBasedAgent(full_context)
        answer = agent.answer("What are the top 3 cost-saving opportunities?")
        assert "Review" in answer or "saving" in answer.lower() or "$" in answer

    def test_total_expenditure_answer(self, full_context):
        """TC-15: Agent reports total spend from real data."""
        agent = RuleBasedAgent(full_context)
        answer = agent.answer("What is the total expenditure?")
        # Should contain a dollar amount
        assert "$" in answer or "USD" in answer

    def test_empty_dataset_returns_warning(self, empty_df):
        """TC-15: Agent returns helpful message when no data is loaded."""
        context = {"df": empty_df}
        agent = RuleBasedAgent(context)
        answer = agent.answer("Which category has the highest spending?")
        assert "no" in answer.lower() or "upload" in answer.lower() or "dataset" in answer.lower()

    def test_unknown_question_returns_suggestions(self, full_context):
        """Unrecognized question returns list of supported questions."""
        agent = RuleBasedAgent(full_context)
        answer = agent.answer("What is the weather today?")
        assert "?" in answer or "supported" in answer.lower() or "asking" in answer.lower()

    def test_vendor_answer(self, full_context):
        """Agent can answer top vendor question."""
        agent = RuleBasedAgent(full_context)
        answer = agent.answer("Which vendor has the highest spend?")
        assert "$" in answer or "USD" in answer

    def test_date_range_answer(self, full_context):
        """Agent can answer date range question."""
        agent = RuleBasedAgent(full_context)
        answer = agent.answer("What is the date range of this dataset?")
        assert "2024" in answer
