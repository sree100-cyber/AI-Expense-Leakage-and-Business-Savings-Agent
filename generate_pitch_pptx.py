"""
generate_pitch_pptx.py
Generates a modern, professional PowerPoint presentation (.pptx)
for the AI Expense Leakage and Business Savings Agent.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Paths
DESKTOP_PATH = r"C:\Users\SREE\OneDrive\Desktop\AI_Expense_Leakage_Agent_Presentation.pptx"
REPO_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "presentation.pptx")

# Initialize presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette
COLOR_BG = RGBColor(15, 23, 42)          # #0F172A
COLOR_CARD = RGBColor(30, 41, 59)        # #1E293B
COLOR_CARD_BORDER = RGBColor(51, 65, 85) # #334155
COLOR_PRIMARY = RGBColor(37, 99, 235)    # #2563EB
COLOR_ACCENT = RGBColor(16, 185, 129)    # #10B981
COLOR_WARN = RGBColor(245, 158, 11)      # #F59E0B
COLOR_WHITE = RGBColor(248, 250, 252)    # #F8FAFC
COLOR_MUTED = RGBColor(148, 163, 184)    # #94A3B8

def set_slide_background(slide):
    """Draw a dark background covering the entire slide."""
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_BG
    bg.line.color.rgb = COLOR_BG

def add_header(slide, tag_text, title_text, subtitle_text=None):
    """Add a unified header with category tag and main title."""
    # Tag
    tb_tag = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
    tf_tag = tb_tag.text_frame
    tf_tag.word_wrap = True
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = tag_text.upper()
    p_tag.font.name = "Arial"
    p_tag.font.size = Pt(11)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_ACCENT

    # Title
    tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.8))
    tf_title = tb_title.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = "Arial"
    p_title.font.size = Pt(28)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_WHITE

    if subtitle_text:
        tb_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.55), Inches(11.7), Inches(0.5))
        tf_sub = tb_sub.text_frame
        tf_sub.word_wrap = True
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = subtitle_text
        p_sub.font.name = "Arial"
        p_sub.font.size = Pt(13)
        p_sub.font.color.rgb = COLOR_MUTED

def add_card(slide, left, top, width, height, title, body_bullets, icon=""):
    """Draw a rounded dark card with title and bullet points."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_CARD
    card.line.color.rgb = COLOR_CARD_BORDER
    card.line.width = Pt(1)

    tb = slide.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.2), Inches(width - 0.5), Inches(height - 0.4))
    tf = tb.text_frame
    tf.word_wrap = True

    # Card Title
    p_title = tf.paragraphs[0]
    p_title.text = f"{icon} {title}".strip()
    p_title.font.name = "Arial"
    p_title.font.size = Pt(16)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_WHITE

    # Bullets
    for b in body_bullets:
        p = tf.add_paragraph()
        p.text = f"•  {b}"
        p.font.name = "Arial"
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_MUTED
        p.space_before = Pt(6)

# ==============================================================================
# SLIDE 1: Title Slide
# ==============================================================================
slide1 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide1)

# Badge
badge = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.2), Inches(3.2), Inches(0.4))
badge.fill.solid()
badge.fill.fore_color.rgb = RGBColor(30, 58, 138)
badge.line.color.rgb = COLOR_PRIMARY
tb_badge = slide1.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(3.2), Inches(0.4))
p_badge = tb_badge.text_frame.paragraphs[0]
p_badge.text = "AI FINANCIAL INTELLIGENCE"
p_badge.font.name = "Arial"
p_badge.font.size = Pt(11)
p_badge.font.bold = True
p_badge.font.color.rgb = COLOR_WHITE
p_badge.alignment = PP_ALIGN.CENTER

# Main Title
tb_t1 = slide1.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(2.0))
tf_t1 = tb_t1.text_frame
tf_t1.word_wrap = True
p_t1 = tf_t1.paragraphs[0]
p_t1.text = "AI Expense Leakage & Business Savings Agent"
p_t1.font.name = "Arial"
p_t1.font.size = Pt(40)
p_t1.font.bold = True
p_t1.font.color.rgb = COLOR_WHITE

p_t1_sub = tf_t1.add_paragraph()
p_t1_sub.text = "Autonomous financial auditing for small businesses & startups. Detects duplicate payments, statistical anomalies, recurring subscription creep, and generates auditable savings plans."
p_t1_sub.font.name = "Arial"
p_t1_sub.font.size = Pt(16)
p_t1_sub.font.color.rgb = COLOR_MUTED
p_t1_sub.space_before = Pt(12)

# Author Info
tb_author = slide1.shapes.add_textbox(Inches(0.8), Inches(4.2), Inches(6.0), Inches(0.8))
p_author = tb_author.text_frame.paragraphs[0]
p_author.text = "Developer & Architect: THAISREE S\nTechnology: Python 3.11 | Streamlit | Scikit-Learn | SQLite3"
p_author.font.name = "Arial"
p_author.font.size = Pt(13)
p_author.font.color.rgb = COLOR_PRIMARY

# 3 Stat Cards on Title Slide
add_card(slide1, 0.8, 5.2, 3.6, 1.6, "100% Non-Invasive", ["No bank credentials required", "In-memory CSV/Excel analysis", "Enterprise privacy compliant"], "🛡️")
add_card(slide1, 4.8, 5.2, 3.6, 1.6, "49/49 Tests Passed", ["100% pytest test suite pass", "All 15 mandatory test cases", "Deterministic reliability"], "🧪")
add_card(slide1, 8.8, 5.2, 3.6, 1.6, "Zero Hallucinations", ["Factual rule & ML engines", "Never invents metrics or spend", "Transparent calculation formulas"], "🎯")


# ==============================================================================
# SLIDE 2: Problem Statement
# ==============================================================================
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide2)
add_header(slide2, "The Business Challenge", "The $300B Silent Bleed in Corporate Expenses", "Small businesses and startups lose 5% to 8% of operating budget annually to preventable financial leakage.")

add_card(slide2, 0.8, 2.3, 3.6, 4.5, "Duplicate Payments", [
    "Dual payments from identical vendor references.",
    "Re-submitted invoices across close date windows.",
    "Accounts payable teams lack automated cross-referencing tools.",
    "Result: Substantial capital lost in unrecovered double billing."
], "🔁")

add_card(slide2, 4.8, 2.3, 3.6, 4.5, "Creeping Subscriptions", [
    "Unmonitored SaaS tools and recurring subscriptions.",
    "Forgotten developer/marketing seat licenses.",
    "Zero ongoing usage auditing across remote teams.",
    "Result: Compounding recurring expenses draining monthly cash flow."
], "📈")

add_card(slide2, 8.8, 2.3, 3.6, 4.5, "Untracked Spending Surges", [
    "Sudden spending spikes (ad spend, equipment, consulting).",
    "Outlier expenses that deviate from historical baselines.",
    "Detected only after quarterly book reconciliation.",
    "Result: Budget overruns before management can intervene."
], "⚡")


# ==============================================================================
# SLIDE 3: The Solution
# ==============================================================================
slide3 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide3)
add_header(slide3, "The AI Solution", "Autonomous, Explainable Financial Guardrails", "An AI-powered agent delivering immediate financial clarity, explainable flags, and measurable savings.")

add_card(slide3, 0.8, 2.3, 5.6, 2.1, "Privacy-First Data Ingestion", [
    "Flexible column alias mapping (supports arbitrary column naming).",
    "Rigorous data validation (rejects invalid dates & amounts).",
    "Zero permanent storage without explicit user export."
], "📂")

add_card(slide3, 6.8, 2.3, 5.6, 2.1, "Hybrid Detection Engine", [
    "Deterministic rules for exact & proximity duplicate payments.",
    "Z-score statistical bounds across vendors & categories.",
    "Unsupervised Scikit-Learn Isolation Forest machine learning."
], "🧠")

add_card(slide3, 0.8, 4.7, 5.6, 2.1, "Assumption-Based Savings", [
    "Conservative, transparent savings calculation formulas.",
    "De-duplication ensures no savings opportunity is counted twice.",
    "Interactive decision matrix: Pending, Accepted, Rejected, Implemented."
], "💡")

add_card(slide3, 6.8, 4.7, 5.6, 2.1, "Factual Conversational Assistant", [
    "Answers complex spending questions strictly using computed data.",
    "Never invents figures or fabricates insights.",
    "Runs fully offline with rule-based fallback or optional LLM integration."
], "🤖")


# ==============================================================================
# SLIDE 4: Architecture & Tech Stack
# ==============================================================================
slide4 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide4)
add_header(slide4, "System Architecture", "Modular Multi-Layer Pipeline", "Designed with strict separation of concerns across 10 specialized Python modules.")

add_card(slide4, 0.8, 2.3, 3.6, 4.5, "Ingestion & Validation", [
    "modules/data_loader.py: Ingests CSV/Excel files; auto-maps canonical aliases.",
    "modules/validation.py: Sanitizes dates, caps bounds, and splits clean vs rejected records.",
    "Displays acceptance rates & error tables in real-time."
], "📥")

add_card(slide4, 4.8, 2.3, 3.6, 4.5, "Detection & Analytics", [
    "modules/duplicate_detection.py: Exact reference & 3-day proximity matching.",
    "modules/anomaly_detection.py: Z-score + Isolation Forest model.",
    "modules/expense_trends.py: Month-over-month category surge tracking."
], "⚙️")

add_card(slide4, 8.8, 2.3, 3.6, 4.5, "Reporting & Persistence", [
    "modules/recommendation_engine.py: Deduplicated savings calculations.",
    "modules/database.py: Relational SQLite persistence with WAL mode.",
    "modules/report_generator.py: Consolidated 8-sheet Excel workbook."
], "📊")


# ==============================================================================
# SLIDE 5: Detection Algorithms
# ==============================================================================
slide5 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide5)
add_header(slide5, "Algorithmic Rigor", "Evidence-Based Detection Modules", "Every flag includes documented justification, contextual baselines, and confidence levels.")

add_card(slide5, 0.8, 2.3, 5.6, 2.1, "Exact & Proximity Duplicates", [
    "Exact Match: Identical non-null payment references (High Confidence).",
    "Proximity Match: Same vendor + amount within ±3 days tolerance window.",
    "Distinguishes exact from potential duplicates without deleting records."
], "👥")

add_card(slide5, 6.8, 2.3, 5.6, 2.1, "Isolation Forest + Z-Score Outliers", [
    "Z-score > 2.5 flags deviations globally and within peer categories.",
    "Isolation Forest (200 trees, 5% contamination) learns multi-dimensional feature patterns.",
    "Degrades gracefully to pure statistical mode if dataset < 20 rows."
], "🌲")

add_card(slide5, 0.8, 4.7, 5.6, 2.1, "Spending Surge Detection", [
    "Aggregates monthly period spend by category, vendor, and department.",
    "Flags month-over-month surges >= 25% with explicit dollar increases.",
    "Pinpoints cost drivers before they compound across quarters."
], "📈")

add_card(slide5, 6.8, 4.7, 5.6, 2.1, "Recurring Subscription Classifier", [
    "Scans multi-month vendor occurrences across the calendar year.",
    "Coefficient of Variation (std/mean < 0.15) identifies fixed SaaS tiers.",
    "Highlights duplicate licenses and unreviewed recurring overhead."
], "🔄")


# ==============================================================================
# SLIDE 6: Savings Engine & Decision Matrix
# ==============================================================================
slide6 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide6)
add_header(slide6, "Financial Impact", "Assumption-Based Cost Optimization Engine", "Conservative, transparent estimates preventing double-counting while maintaining executive audit trails.")

add_card(slide6, 0.8, 2.3, 3.6, 4.5, "1. Conservative Assumptions", [
    "Duplicate Payments: 8% monthly recovery potential.",
    "Outlier Expenses: 5% budget recovery through verification.",
    "Category Surges: 15% optimization above baseline trend.",
    "Subscriptions: 10% savings via license tier rightsizing.",
    "Annual savings strictly = Monthly × 12."
], "🎯")

add_card(slide6, 4.8, 2.3, 3.6, 4.5, "2. Strict De-Duplication", [
    "Recommendation de-duplication eliminates double-counting.",
    "Distinct category/vendor keys prevent overlapping estimates.",
    "Distinguishes Potential Savings from Realized Savings.",
    "Auditable methodology backed by verifiable evidence."
], "🛡️")

add_card(slide6, 8.8, 2.3, 3.6, 4.5, "3. SQLite Decision Matrix", [
    "Finance managers assign statuses: Pending, Accepted, Rejected, Implemented.",
    "Decisions & review notes persist across application reloads.",
    "Relational SQLite schema tracks historical analysis sessions.",
    "Full governance audit trail for leadership reviews."
], "📋")


# ==============================================================================
# SLIDE 7: Automated Test Suite (49/49 Pass)
# ==============================================================================
slide7 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide7)
add_header(slide7, "Software Quality Assurance", "Comprehensive Pytest Test Suite — 49/49 Passed", "Verified against 49 rigorous unit & integration tests covering all 15 required test cases in 2.54s.")

add_card(slide7, 0.8, 2.3, 5.6, 4.5, "Test Execution Summary", [
    "tests/test_duplicates.py: 9/9 PASSED (TC-01, TC-02, TC-03)",
    "tests/test_anomalies.py: 7/7 PASSED (TC-04, TC-10)",
    "tests/test_data_validation.py: 11/11 PASSED (TC-05, TC-06, TC-11, TC-12)",
    "tests/test_recommendations.py: 14/14 PASSED (TC-07, TC-08, TC-09, TC-13)",
    "tests/test_reports.py: 8/8 PASSED (TC-14, TC-15)",
    "Total: 49 passing tests in 2.54 seconds",
    "Zero fabricated results — real automated test log saved in repo."
], "🧪")

add_card(slide7, 6.8, 2.3, 5.6, 4.5, "Key Criteria Validated", [
    "✅ Exact duplicate payment detected by payment reference.",
    "✅ Possible duplicate with minor date difference flagged.",
    "✅ Normal transactions not falsely classified as duplicates.",
    "✅ Unusually large transactions flagged under statistical rules.",
    "✅ Missing/invalid amounts and corrupted dates safely handled.",
    "✅ Annual savings strictly uses monthly × 12 multiplier.",
    "✅ De-duplication guarantees no double-counting.",
    "✅ Recommendation status changes persist in SQLite database.",
    "✅ Rule-based assistant answers queries using real computed data."
], "✅")


# ==============================================================================
# SLIDE 8: Deployment & Conclusion
# ==============================================================================
slide8 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide8)
add_header(slide8, "Production Readiness", "Deployment Ready & Future Roadmap", "Complete working codebase, reproducible sample dataset, and deployment-ready architecture.")

add_card(slide8, 0.8, 2.3, 5.6, 4.5, "Production Deliverables", [
    "GitHub Repository: https://github.com/sree100-cyber/AI-Expense-Leakage-and-Business-Savings-Agent",
    "Interactive UI: 8 Streamlit dashboard pages with Plotly visualizations.",
    "Sample Dataset: 157 synthetic transactions spanning 12 months with planted test leaks.",
    "Standalone Test Upload: test_upload_expenses.csv & .xlsx for immediate testing.",
    "One-Click Deploy: Fully prepared for Streamlit Community Cloud."
], "🚀")

add_card(slide8, 6.8, 2.3, 5.6, 4.5, "Future Development Roadmap", [
    "Direct Accounting API Sync: QuickBooks, Xero, and NetSuite live connectors.",
    "OCR Receipt & Invoice Processing: Extract line items from scanned paper receipts.",
    "Real-Time Slack / Teams Alerts: Notify finance teams when budget anomalies trigger.",
    "Multi-Currency Auto-Conversion: Live FX exchange rates across global subsidiaries.",
    "Enterprise Role-Based Access: Departmental budget limits & approval hierarchies."
], "🔮")

# Save Presentation
prs.save(DESKTOP_PATH)
prs.save(REPO_PATH)

print(f"[OK] PowerPoint Presentation created successfully!")
print(f"     Desktop: {DESKTOP_PATH}")
print(f"     Repo:    {REPO_PATH}")
